import argparse
import json
import os
import random
import re
import sqlite3
import sys
import time
from datetime import datetime

import openpyxl
from playwright.sync_api import sync_playwright

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(line_buffering=True)

import resume_tailor
import vault_manager
import cover_letter_generator

DEFAULT_EXCEL_FILE = "JobHive_Jobs.xlsx"
DEFAULT_DB_FILE = "jobs.sqlite"
DEFAULT_VAULT_DIR = "." if (os.path.isdir("Jobs") or os.path.exists("00-Dashboard.md")) else "JobHiveVault"

def load_jobs_from_excel(excel_path=DEFAULT_EXCEL_FILE):
    """Read all jobs and their direct URLs from JobHive_Jobs.xlsx."""
    if not os.path.exists(excel_path):
        print(f"[!] Excel file not found: {excel_path}")
        return []

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb.active
    jobs = []

    for r in range(2, ws.max_row + 1):
        date_found = ws.cell(r, 1).value
        title = ws.cell(r, 2).value
        company = ws.cell(r, 3).value
        source = ws.cell(r, 4).value
        url = ws.cell(r, 5).value
        resume_name = ws.cell(r, 6).value
        cover_letter = ws.cell(r, 7).value
        status = ws.cell(r, 8).value if ws.cell(r, 8).value else (ws.cell(r, 7).value if ws.cell(r, 7).value in ["Applied", "Interview", "Decline"] else "")

        if not url:
            continue

        # Extract Job ID from URL or resume name
        job_id = None
        m = re.search(r'/job/(\d+)', str(url))
        if m:
            job_id = m.group(1)
        elif resume_name:
            m_res = re.search(r'(\d+)_resume', str(resume_name))
            if m_res:
                job_id = m_res.group(1)

        if not job_id:
            continue

        jobs.append({
            "row": r,
            "job_id": str(job_id),
            "title": str(title or "Untitled Job").strip(),
            "company": str(company or "Company").strip(),
            "url": str(url).strip(),
            "date_found": str(date_found or datetime.now().strftime("%Y-%m-%d")),
            "status": str(status or "").strip()
        })

    return jobs

def fetch_full_job_description(page, url):
    """Extract full job description from JobStreet posting page."""
    try:
        page.goto(url, wait_until="domcontentloaded", timeout=25000)
        # Jitter pause to allow DOM hydration
        page.wait_for_timeout(random.randint(1800, 2600))

        if "Just a moment" in page.title():
            print("    [!] Cloudflare challenge detected. Waiting 5s...")
            page.wait_for_timeout(5000)

        # Main JobStreet selector for job details
        el = page.query_selector("[data-automation='jobAdDetails']")
        if not el:
            el = page.query_selector("[data-automation='jobDescription']")
        if not el:
            el = page.query_selector("div[class*='jobDetails']")

        if el:
            text = el.inner_text().strip()
            if text:
                return text

        # Fallback: grab all paragraphs if specific container missing
        main_body = page.query_selector("main") or page.query_selector("article")
        if main_body:
            return main_body.inner_text().strip()

        return None
    except Exception as e:
        print(f"    [!] Error loading {url}: {e}")
        return None

def update_job_in_db(db_path, job_id, full_desc):
    """Store scraped full description into SQLite."""
    if not os.path.exists(db_path):
        return
    conn = sqlite3.connect(db_path)
    cur = conn.cursor()
    cur.execute("UPDATE jobs SET full_description = ? WHERE job_id = ?", (full_desc, job_id))
    conn.commit()
    conn.close()

def run_recrawl_and_retailor(excel_path=DEFAULT_EXCEL_FILE, db_path=DEFAULT_DB_FILE, vault_dir=DEFAULT_VAULT_DIR, limit=None, target_job_id=None):
    """
    1. Loads all job links from JobHive_Jobs.xlsx.
    2. Uses Playwright to crawl each job's full description.
    3. Remakes the tailored resume using the enhanced multi-section tailoring engine.
    4. Updates Obsidian job notes and refreshes dashboard & Excel tracker.
    """
    jobs = load_jobs_from_excel(excel_path)
    if not jobs:
        print("[!] No jobs found in Excel. Aborting.")
        return

    if target_job_id:
        jobs = [j for j in jobs if j["job_id"] == str(target_job_id)]
        if not jobs:
            print(f"[!] Job ID {target_job_id} not found in Excel.")
            return

    if limit and limit > 0:
        jobs = jobs[:limit]

    print("\n" + "=" * 60)
    print(f"[*] Starting Link-Based Recrawl & Resume Remake Pipeline")
    print(f"[*] Total links to process: {len(jobs)}")
    print(f"[*] Source Excel: {excel_path}")
    print(f"[*] Obsidian Vault: {vault_dir}")
    print("=" * 60)

    success_count = 0
    scraped_count = 0

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        ctx = browser.new_context(
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
            viewport={"width": 1280, "height": 800}
        )
        page = ctx.new_page()

        for idx, job in enumerate(jobs, start=1):
            jid = job["job_id"]
            title = job["title"]
            company = job["company"]
            url = job["url"]

            print(f"\n>>> [{idx}/{len(jobs)}] Job {jid}: {title} @ {company}")
            print(f"    URL: {url}")

            # 1. Scrape full job description
            full_desc = fetch_full_job_description(page, url)
            if full_desc:
                scraped_count += 1
                job["full_description"] = full_desc
                update_job_in_db(db_path, jid, full_desc)
                print(f"    [+] Scraped Full Description ({len(full_desc)} characters)")
            else:
                print("    [!] Full description extraction empty. Using existing card summary.")
                # Fallback to existing note / SQLite
                job["full_description"] = None

            # 2. Invoke upgraded Resume Tailor
            print(f"    [*] Generating Deeply Tailored Resume...")
            ok, err = resume_tailor.tailor_job(vault_dir, job)
            if ok:
                success_count += 1
                print(f"    [OK] Remade Resume for Job {jid}!")
            else:
                print(f"    [!] Tailor error: {err}")

            # 3. Generate Tailored Cover Letter
            print(f"    [*] Generating Tailored Cover Letter...")
            cl_ok, cl_err = cover_letter_generator.create_cover_letter_for_job(vault_dir, job)
            if cl_ok:
                print(f"    [OK] Generated Cover Letter for Job {jid}!")
            else:
                print(f"    [!] Cover letter error: {cl_err}")

            # 4. Inter-job polite jitter delay (2.0s - 3.5s) to avoid bot detection
            if idx < len(jobs):
                jitter = random.uniform(2.0, 3.5)
                time.sleep(jitter)

        browser.close()

    # Rebuild Obsidian Dashboard and refresh Excel Tracker
    print("\n[*] Rebuilding Dashboard & Excel Tracker...")
    vault_manager.regenerate_dashboard(vault_dir)
    print(f"[+] Complete! Successfully remade {success_count}/{len(jobs)} resumes.")
    print("=" * 60)

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Recrawl job descriptions from Excel links and remake resumes")
    parser.add_argument("--excel", default=DEFAULT_EXCEL_FILE, help="Path to JobHive_Jobs.xlsx")
    parser.add_argument("--db", default=DEFAULT_DB_FILE, help="Path to jobs.sqlite")
    parser.add_argument("--vault", default=DEFAULT_VAULT_DIR, help="Path to vault")
    parser.add_argument("--limit", type=int, default=None, help="Limit number of jobs to crawl (e.g. for testing)")
    parser.add_argument("--job-id", help="Crawl a single specific job ID")
    args = parser.parse_args()

    run_recrawl_and_retailor(
        excel_path=args.excel,
        db_path=args.db,
        vault_dir=args.vault,
        limit=args.limit,
        target_job_id=args.job_id
    )
