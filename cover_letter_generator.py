import argparse
import json
import os
import re
import sys
import urllib.request
from datetime import datetime

import docx
from docx.shared import Inches, Pt, RGBColor

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(line_buffering=True)

import resume_tailor
import vault_manager

COVER_LETTER_DIR_REL = os.path.join("Resumes", "Cover_Letters")

def generate_cover_letter_content(base_resume_text, job_title, company, job_description, match_notes=""):
    """
    Calls LLM to generate a customized, authentic cover letter.
    Enforces strict zero-fabrication rules.
    """
    cfg = resume_tailor.load_llm_config()
    base_url = cfg.get("base_url", "https://api.xkiro.com/v1").rstrip("/")
    api_key = cfg.get("api_key", "")
    model = cfg.get("model", "qwen/qwen3.7-flash:free")

    system_prompt = (
        "You are an expert technical cover letter writer and hiring advisor.\n"
        "Your task is to write a highly compelling, authentic, professional cover letter "
        "tailored specifically to the provided job posting, using ONLY the candidate's real qualifications.\n\n"
        "STRICT GUARDRAILS:\n"
        "1. NEVER invent, fabricate, or exaggerate experience, companies, degrees, or skills not present in the base resume.\n"
        "2. If the posting mentions tools or frameworks the candidate lacks (e.g., n8n, Make.com), do not claim mastery. "
        "Instead, frame it through learning agility and strong transferable foundation (e.g. REST APIs, webhooks, scripting).\n"
        "3. Avoid generic AI fluff and boilerplate clichés (e.g., 'I am thrilled to submit my resume', 'In today's fast-paced world', 'dynamic synergy'). "
        "Keep the tone direct, engineering-focused, authentic, and confident.\n"
        "4. Address any explicit job constraints (e.g. U.S. business hours, remote setup, English proficiency, tech allowance) naturally.\n"
        "5. Output must be valid JSON with no markdown wrapping, conforming to the requested schema."
    )

    user_prompt = f"""Target Job Title: {job_title}
Target Company: {company}

Target Job Description:
{job_description}

Match Notes / Analysis:
{match_notes}

Candidate Base Resume:
{base_resume_text}

Provide your response in strictly valid JSON with this exact schema:
{{
  "salutation": "Dear Hiring Team, (or Dear [Company] Team,)",
  "opening_paragraph": "Strong opening stating position, immediate availability for US/EST hours or remote arrangement, and high-level fit.",
  "body_paragraph_1": "Focus on candidate's direct technical accomplishments (e.g. specific projects like CellWatch for AI/Python, eKalinga+ for C#/data sync, or Car-Showroom/eKalingaWeb for Laravel/PHP) that prove competence.",
  "body_paragraph_2": "Focus on APIs, integrations, authentication, webhooks, or architecture relevant to the posting, acknowledging enthusiasm for learning company-specific tools.",
  "body_paragraph_3": "Focus on communication, discipline, dedicated home office setup, reliability, and work ethic.",
  "closing_paragraph": "Professional call to action expressing enthusiasm for interview/discussion.",
  "sign_off": "Sincerely,"
}}"""

    url = f"{base_url}/chat/completions"
    payload = {
        "model": model,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt}
        ],
        "response_format": {"type": "json_object"}
    }
    headers = {
        "Content-Type": "application/json",
        "Authorization": f"Bearer {api_key}",
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
    }

    try:
        req = urllib.request.Request(url, data=json.dumps(payload).encode("utf-8"), headers=headers, method="POST")
        with urllib.request.urlopen(req, timeout=60) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            raw_text = data["choices"][0]["message"]["content"]
    except Exception as e:
        return None, f"LLM API call failed: {e}"

    cleaned = re.sub(r'^```json\s*', '', raw_text.strip(), flags=re.IGNORECASE)
    cleaned = re.sub(r'\s*```$', '', cleaned.strip())

    try:
        parsed = json.loads(cleaned)
        return parsed, None
    except json.JSONDecodeError as err:
        return None, f"JSON parse error: {err}. Raw: {raw_text[:200]}"

def build_cover_letter_docx(output_path, letter_data, job_title, company, job_id):
    """Generates an executive-styled Word document (.docx) cover letter."""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc = docx.Document()

    # 1-inch margins
    for sec in doc.sections:
        sec.top_margin = Inches(1)
        sec.bottom_margin = Inches(1)
        sec.left_margin = Inches(1)
        sec.right_margin = Inches(1)

    # Name Header (Navy)
    p_name = doc.add_paragraph()
    p_name.paragraph_format.space_after = Pt(2)
    r_name = p_name.add_run("BIEN JOSEF G. REGIDOR")
    r_name.bold = True
    r_name.font.size = Pt(16)
    r_name.font.name = "Calibri"
    r_name.font.color.rgb = RGBColor(27, 78, 121)

    # Contact line
    p_contact = doc.add_paragraph()
    p_contact.paragraph_format.space_after = Pt(16)
    r_contact = p_contact.add_run("Full-Stack & Systems Developer  |  Digos City, Philippines  |  benregidor@gmail.com  |  github.com/aze-last")
    r_contact.font.size = Pt(9.5)
    r_contact.font.name = "Calibri"
    r_contact.font.color.rgb = RGBColor(89, 89, 89)

    # Date & Recipient
    today_str = datetime.now().strftime("%B %d, %Y")
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(12)
    r_meta = p_meta.add_run(f"{today_str}\n\nHiring Team\n{company}\nRe: Application for {job_title} (Job ID: {job_id})")
    r_meta.font.size = Pt(10)
    r_meta.font.name = "Calibri"

    # Salutation
    salutation = letter_data.get("salutation", "Dear Hiring Team,")
    p_sal = doc.add_paragraph(salutation)
    p_sal.paragraph_format.space_after = Pt(10)
    for r in p_sal.runs:
        r.font.size = Pt(10)
        r.font.name = "Calibri"

    # Body paragraphs
    paragraphs = [
        letter_data.get("opening_paragraph", ""),
        letter_data.get("body_paragraph_1", ""),
        letter_data.get("body_paragraph_2", ""),
        letter_data.get("body_paragraph_3", ""),
        letter_data.get("closing_paragraph", "")
    ]

    for text in paragraphs:
        if not text:
            continue
        p = doc.add_paragraph(text)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = 1.15
        for r in p.runs:
            r.font.size = Pt(10)
            r.font.name = "Calibri"

    # Sign-off
    sign_off = letter_data.get("sign_off", "Sincerely,")
    p_close = doc.add_paragraph(f"{sign_off}\n\nBien Josef G. Regidor")
    p_close.paragraph_format.space_after = Pt(0)
    for r in p_close.runs:
        r.font.size = Pt(10)
        r.font.name = "Calibri"

    try:
        doc.save(output_path)
    except PermissionError:
        alt_path = output_path.replace(".docx", "_new.docx")
        try:
            doc.save(alt_path)
            print(f"  [+] Word file locked; saved to: {alt_path}")
        except Exception:
            pass

    return output_path

def build_cover_letter_markdown(output_path, letter_data, job_title, company, job_id, docx_rel):
    """Generates a clean Markdown (.md) cover letter document."""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    today_str = datetime.now().strftime("%B %d, %Y")

    content = f"""# Cover Letter: {job_title}

**Applicant:** Bien Josef G. Regidor  
**Contact:** benregidor@gmail.com | Digos City, Davao del Sur, Philippines | [github.com/aze-last](https://github.com/aze-last)  
**Target Role:** {job_title}  
**Employer:** {company} (Job ID: `{job_id}`)  
**Word Document:** [`{docx_rel}`](file:///{os.path.abspath(docx_rel).replace('\\', '/')})

---

**{today_str}**

**Hiring Team**  
{company}  
*Re: Application for {job_title} (Job ID: {job_id})*

{letter_data.get('salutation', 'Dear Hiring Team,')}

{letter_data.get('opening_paragraph', '')}

{letter_data.get('body_paragraph_1', '')}

{letter_data.get('body_paragraph_2', '')}

{letter_data.get('body_paragraph_3', '')}

{letter_data.get('closing_paragraph', '')}

{letter_data.get('sign_off', 'Sincerely,')}  
**Bien Josef G. Regidor**  
[github.com/aze-last](https://github.com/aze-last) • benregidor@gmail.com
"""

    with open(output_path, "w", encoding="utf-8") as f:
        f.write(content)

    return output_path

def create_cover_letter_for_job(vault_dir, job_data):
    """
    End-to-end cover letter generation for a specific job:
    1. Reads base resume.
    2. Calls LLM with role-specific prompt.
    3. Generates Word (.docx) and Markdown (.md).
    4. Updates Obsidian job note with link.
    """
    job_id = str(job_data.get("job_id", ""))
    title = job_data.get("title", "Untitled Job")
    company = job_data.get("company", "Company")
    desc = job_data.get("full_description") or job_data.get("summary") or ""
    match_notes = job_data.get("match_notes", "")

    if not job_id:
        return False, "Missing job_id"

    base_docx = os.path.join(vault_dir, "Resumes", "Base", "resume_base.docx")
    base_text, err = resume_tailor.extract_base_resume_text(base_docx)
    if err:
        return False, err

    print(f"\n[Cover Letter Subagent] Generating Cover Letter for Job {job_id}: {title} @ {company}")

    # Call LLM
    letter_data, llm_err = generate_cover_letter_content(base_text, title, company, desc, match_notes)
    if llm_err:
        print(f"  [!] Cover letter generation failed: {llm_err}")
        return False, llm_err

    # File paths
    docx_rel = f"Resumes/Cover_Letters/{job_id}_cover_letter.docx"
    docx_full = os.path.join(vault_dir, docx_rel)
    md_rel = f"Resumes/Cover_Letters/{job_id}_cover_letter.md"
    md_full = os.path.join(vault_dir, md_rel)

    # Build Word & Markdown
    build_cover_letter_docx(docx_full, letter_data, title, company, job_id)
    build_cover_letter_markdown(md_full, letter_data, title, company, job_id, docx_rel)

    print(f"  [+] Word Cover Letter: {docx_rel}")
    print(f"  [+] Markdown Cover Letter: {md_rel}")

    # Update Obsidian Note
    note_path = vault_manager.get_job_note_path(vault_dir, job_id, title)
    if os.path.exists(note_path):
        with open(note_path, "r", encoding="utf-8") as f:
            content = f.read()

        meta, body = vault_manager.parse_frontmatter(content)
        meta["cover_letter"] = docx_rel

        # Ensure cover letter link is in body
        cl_link_line = f"- **Cover Letter:** [`{docx_rel}`](file:///{os.path.abspath(docx_full).replace('\\', '/')}) | [[{md_rel}|View Markdown]]"
        
        # Check if already in body
        if "- **Cover Letter:**" not in body:
            if "- **Tailored Resume:**" in body:
                body = body.replace(
                    "- **Tailored Resume:**",
                    f"- **Tailored Resume:**"
                )
                # insert after tailored resume line
                lines = body.splitlines()
                new_lines = []
                for line in lines:
                    new_lines.append(line)
                    if "- **Tailored Resume:**" in line:
                        new_lines.append(cl_link_line)
                body = "\n".join(new_lines)
            else:
                body = cl_link_line + "\n\n" + body

        # Rebuild frontmatter
        fm_lines = ["---"]
        for k, v in meta.items():
            if isinstance(v, bool):
                fm_lines.append(f"{k}: {str(v).lower()}")
            else:
                fm_lines.append(f"{k}: \"{v}\"")
        fm_lines.append("---")
        updated_note = "\n".join(fm_lines) + "\n\n" + body

        with open(note_path, "w", encoding="utf-8") as f:
            f.write(updated_note)
        print(f"  [+] Updated Job Note: {os.path.basename(note_path)}")

    return True, None

def batch_generate_cover_letters(vault_dir=".", skip_existing=True, limit=None, max_workers=3):
    """
    Scans all job notes in vault and generates cover letters for them concurrently.
    If skip_existing is True, skips jobs that already have a generated cover letter.
    """
    import concurrent.futures
    import threading

    jobs_dir = os.path.join(vault_dir, "Jobs")
    if not os.path.exists(jobs_dir):
        print(f"[!] Jobs directory not found at {jobs_dir}")
        return 0

    all_notes = [f for f in os.listdir(jobs_dir) if f.endswith(".md")]
    print(f"[*] Found {len(all_notes)} job notes in vault.")

    # Collect jobs to process
    jobs_to_process = []
    for fname in all_notes:
        note_path = os.path.join(jobs_dir, fname)
        try:
            with open(note_path, "r", encoding="utf-8") as f:
                content = f.read()
            meta, body = vault_manager.parse_frontmatter(content)
        except Exception as e:
            continue

        jid = str(meta.get("job_id", ""))
        if not jid:
            continue

        cl_path = os.path.join(vault_dir, "Resumes", "Cover_Letters", f"{jid}_cover_letter.docx")
        if skip_existing and os.path.exists(cl_path):
            continue

        desc = ""
        match_notes = ""
        if "## Job Description" in body:
            parts = body.split("## Job Description", 1)
            if "## Match Notes" in parts[1]:
                desc_part, match_part = parts[1].split("## Match Notes", 1)
                desc = desc_part.strip()
                match_notes = match_part.strip()
            else:
                desc = parts[1].strip()

        meta["full_description"] = desc
        meta["match_notes"] = match_notes
        jobs_to_process.append(meta)

    if limit and limit > 0:
        jobs_to_process = jobs_to_process[:limit]

    total_target = len(jobs_to_process)
    print(f"[*] Total cover letters to generate: {total_target} (Workers: {max_workers})")
    if total_target == 0:
        print("[+] All jobs already have cover letters! Nothing to generate.")
        vault_manager.regenerate_dashboard(vault_dir)
        return 0

    lock = threading.Lock()
    completed_count = 0

    def process_single_job(meta):
        nonlocal completed_count
        jid = meta.get("job_id")
        title = meta.get("title", "Job")
        ok, err = create_cover_letter_for_job(vault_dir, meta)
        with lock:
            if ok:
                completed_count += 1
                print(f"  [{completed_count}/{total_target}] Generated Cover Letter for Job {jid}: {title}")
                if completed_count % 10 == 0:
                    try:
                        vault_manager.export_to_excel(vault_dir)
                    except Exception:
                        pass
            else:
                print(f"  [!] Failed for Job {jid}: {err}")
        return ok

    with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as executor:
        list(executor.map(process_single_job, jobs_to_process))

    print(f"\n[+] Batch generation finished. Generated {completed_count}/{total_target} cover letters.")
    vault_manager.regenerate_dashboard(vault_dir)
    return completed_count

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Cover Letter Generator Subagent")
    parser.add_argument("--job-id", help="Single Job ID to generate cover letter for")
    parser.add_argument("--all", action="store_true", help="Generate cover letters for all crawled jobs in vault")
    parser.add_argument("--limit", type=int, default=None, help="Limit number of cover letters to generate in batch")
    parser.add_argument("--force", action="store_true", help="Overwrite existing cover letters even if they exist")
    parser.add_argument("--workers", type=int, default=3, help="Number of concurrent workers for batch generation")
    parser.add_argument("--vault", default="." if (os.path.isdir("Jobs") or os.path.exists("00-Dashboard.md")) else "JobHiveVault", help="Path to vault")
    args = parser.parse_args()

    vault_dir = args.vault
    if vault_dir == "JobHiveVault" and not os.path.isdir("JobHiveVault") and (os.path.isdir("Jobs") or os.path.exists("00-Dashboard.md")):
        vault_dir = "."

    if args.all:
        batch_generate_cover_letters(vault_dir=vault_dir, skip_existing=not args.force, limit=args.limit, max_workers=args.workers)
    elif args.job_id:
        jobs_dir = os.path.join(vault_dir, "Jobs")
        target_note = None
        for f in os.listdir(jobs_dir):
            if f.startswith(f"{args.job_id}-") and f.endswith(".md"):
                target_note = os.path.join(jobs_dir, f)
                break

        if not target_note:
            print(f"[!] Note not found for Job ID {args.job_id}")
            sys.exit(1)

        with open(target_note, "r", encoding="utf-8") as file:
            meta, body = vault_manager.parse_frontmatter(file.read())

        desc = ""
        match_notes = ""
        if "## Job Description" in body:
            parts = body.split("## Job Description", 1)
            if "## Match Notes" in parts[1]:
                desc_part, match_part = parts[1].split("## Match Notes", 1)
                desc = desc_part.strip()
                match_notes = match_part.strip()
            else:
                desc = parts[1].strip()

        meta["full_description"] = desc
        meta["match_notes"] = match_notes

        ok, err = create_cover_letter_for_job(vault_dir, meta)
        if not ok:
            print(f"[!] Failed: {err}")
            sys.exit(1)
        else:
            print("[+] Success!")
    else:
        print("[!] Please specify --job-id <id> or --all")

